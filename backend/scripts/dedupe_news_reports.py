"""Deduplicate syndicated copies of the same news story.

Reads a reports directory of per-article folders (meta.json + optional full.txt)
and writes an Excel workbook. Source files are not modified.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_INPUT = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "data",
    "news",
    "v6",
    "reports - 副本",
)
DEFAULT_OUTPUT = os.path.join(
    os.path.dirname(DEFAULT_INPUT),
    "reports_deduped.xlsx",
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
KEPT_FILL = PatternFill("solid", fgColor="E2EFDA")
REMOVED_FILL = PatternFill("solid", fgColor="FCE4D6")
SHINGLE_SIZE = 8
BODY_JACCARD_THRESHOLD = 0.55
MIN_BODY_CHARS = 400


def normalize_title(title: str) -> str:
    text = (title or "").lower()
    for src, dst in (
        ("\u2019", "'"),
        ("\u2018", "'"),
        ("\u201c", '"'),
        ("\u201d", '"'),
        ("\ufffd", ""),
        ("\u2014", " "),
        ("\u2013", " "),
    ):
        text = text.replace(src, dst)
    text = text.replace("|", " ")
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"\s+-\s+[a-z0-9 .,&']{2,40}$", "", text).strip()
    return text


def article_body(path: str) -> str:
    if not os.path.exists(path):
        return ""
    with open(path, encoding="utf-8", errors="replace") as handle:
        raw = handle.read()
    marker = "URL:"
    if marker in raw:
        after_url = raw.split(marker, 1)[1]
        raw = after_url.split("\n", 1)[1] if "\n" in after_url else ""
    text = raw.lower().replace("\ufffd", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def shingles(text: str) -> set[str]:
    words = text.split()
    if len(words) < SHINGLE_SIZE:
        return set()
    return {
        " ".join(words[index : index + SHINGLE_SIZE])
        for index in range(0, len(words) - SHINGLE_SIZE + 1, 2)
    }


def load_articles(root: str) -> list[dict]:
    articles = []
    for folder in sorted(os.listdir(root)):
        folder_path = os.path.join(root, folder)
        meta_path = os.path.join(folder_path, "meta.json")
        if not os.path.isdir(folder_path) or not os.path.exists(meta_path):
            continue
        with open(meta_path, encoding="utf-8") as handle:
            meta = json.load(handle)
        body = article_body(os.path.join(folder_path, "full.txt"))
        title = (meta.get("title") or "").strip()
        articles.append(
            {
                "id": meta.get("url_hash") or folder,
                "title": title,
                "title_key": normalize_title(title),
                "url": meta.get("url") or "",
                "domain": meta.get("domain") or "",
                "source_name": meta.get("source_name") or "",
                "published": meta.get("published") or "",
                "fetch_status": meta.get("fetch_status") or "",
                "language": meta.get("article_language") or "",
                "body_chars": int(meta.get("body_chars") or 0),
                "mention_count": int(meta.get("mention_count") or 0),
                "byline": meta.get("byline") or "",
                "body": body,
                "shingles": shingles(body) if len(body) >= MIN_BODY_CHARS else set(),
                "body_hash": hashlib.md5(body.encode("utf-8", errors="replace")).hexdigest()
                if body
                else "",
            }
        )
    return articles


class UnionFind:
    def __init__(self, size: int) -> None:
        self.parent = list(range(size))

    def find(self, index: int) -> int:
        while self.parent[index] != index:
            self.parent[index] = self.parent[self.parent[index]]
            index = self.parent[index]
        return index

    def union(self, left: int, right: int) -> bool:
        root_left = self.find(left)
        root_right = self.find(right)
        if root_left == root_right:
            return False
        self.parent[root_right] = root_left
        return True


def cluster_articles(articles: list[dict]) -> list[tuple[list[int], str]]:
    union_find = UnionFind(len(articles))
    link_reason: dict[tuple[int, int], str] = {}

    title_buckets: dict[str, list[int]] = defaultdict(list)
    hash_buckets: dict[str, list[int]] = defaultdict(list)
    for index, article in enumerate(articles):
        if article["title_key"] and article["title_key"] != "wire":
            title_buckets[article["title_key"]].append(index)
        if article["body_hash"]:
            hash_buckets[article["body_hash"]].append(index)

    for indexes in title_buckets.values():
        for other in indexes[1:]:
            if union_find.union(indexes[0], other):
                link_reason[(indexes[0], other)] = "标题相同"

    for indexes in hash_buckets.values():
        for other in indexes[1:]:
            if union_find.union(indexes[0], other):
                link_reason[(indexes[0], other)] = "正文相同"

    comparable = [index for index, article in enumerate(articles) if article["shingles"]]
    for position, left in enumerate(comparable):
        left_shingles = articles[left]["shingles"]
        left_len = len(articles[left]["body"])
        for right in comparable[position + 1 :]:
            if union_find.find(left) == union_find.find(right):
                continue
            right_len = len(articles[right]["body"])
            longer = max(left_len, right_len)
            if abs(left_len - right_len) > longer * 0.45:
                continue
            right_shingles = articles[right]["shingles"]
            union_size = len(left_shingles | right_shingles)
            if not union_size:
                continue
            score = len(left_shingles & right_shingles) / union_size
            if score >= BODY_JACCARD_THRESHOLD and union_find.union(left, right):
                link_reason[(left, right)] = f"正文相似({score:.2f})"

    groups: dict[int, list[int]] = defaultdict(list)
    for index in range(len(articles)):
        groups[union_find.find(index)].append(index)

    group_reason: dict[int, str] = {}
    for root, indexes in groups.items():
        if len(indexes) == 1:
            article = articles[indexes[0]]
            if not article["title"] and not article["body"]:
                group_reason[root] = "无标题无正文，单独保留"
            else:
                group_reason[root] = "无重复"
            continue
        reasons = [
            reason
            for (left, right), reason in link_reason.items()
            if union_find.find(left) == root
        ]
        if any(reason == "标题相同" for reason in reasons):
            group_reason[root] = "标题相同"
        elif any(reason == "正文相同" for reason in reasons):
            group_reason[root] = "正文相同"
        else:
            group_reason[root] = "正文相似"
    return [(indexes, group_reason[root]) for root, indexes in groups.items()]


def choose_representative(articles: list[dict], indexes: list[int]) -> int:
    def rank(index: int) -> tuple:
        article = articles[index]
        return (
            article["fetch_status"] == "ok",
            article["body_chars"],
            article["mention_count"],
            bool(article["published"]),
            bool(article["title"]),
        )

    return max(indexes, key=rank)


def build_rows(articles: list[dict]) -> tuple[list[dict], list[dict], dict]:
    grouped = cluster_articles(articles)
    kept_rows = []
    detail_rows = []
    removed = 0
    duplicate_groups = 0

    ordered_groups = sorted(
        grouped,
        key=lambda item: (
            -len(item[0]),
            articles[choose_representative(articles, item[0])]["title"].lower(),
            articles[item[0][0]]["id"],
        ),
    )

    for group_number, (indexes, reason) in enumerate(ordered_groups, start=1):
        kept_index = choose_representative(articles, indexes)
        kept = articles[kept_index]
        others = [articles[index] for index in indexes if index != kept_index]
        if others:
            duplicate_groups += 1
            removed += len(others)
        removed_domains = sorted({item["domain"] for item in others if item["domain"]})
        kept_rows.append(
            {
                "组号": group_number,
                "保留编号": kept["id"],
                "标题": kept["title"],
                "域名": kept["domain"],
                "链接": kept["url"],
                "发布时间": kept["published"],
                "抓取状态": kept["fetch_status"],
                "语言": kept["language"],
                "正文字数": kept["body_chars"],
                "提及次数": kept["mention_count"],
                "作者": kept["byline"],
                "组内条数": len(indexes),
                "消除条数": len(others),
                "判定依据": reason if others or reason.startswith("无标题") else "无重复",
                "被消除来源": "；".join(removed_domains),
            }
        )
        members = [(kept_index, "保留")] + [
            (index, "消除") for index in indexes if index != kept_index
        ]
        for index, role in members:
            article = articles[index]
            detail_rows.append(
                {
                    "组号": group_number,
                    "处理": role,
                    "编号": article["id"],
                    "标题": article["title"],
                    "域名": article["domain"],
                    "链接": article["url"],
                    "发布时间": article["published"],
                    "抓取状态": article["fetch_status"],
                    "正文字数": article["body_chars"],
                    "提及次数": article["mention_count"],
                    "判定依据": reason if len(indexes) > 1 else kept_rows[-1]["判定依据"],
                }
            )

    summary = {
        "原始条数": len(articles),
        "去重后条数": len(kept_rows),
        "消除条数": removed,
        "重复组数": duplicate_groups,
        "无正文单独保留": sum(1 for row in kept_rows if row["判定依据"].startswith("无标题")),
    }
    return kept_rows, detail_rows, summary


def write_sheet(workbook: Workbook, title: str, rows: list[dict], highlight_role: bool = False) -> None:
    sheet = workbook.create_sheet(title)
    if not rows:
        sheet["A1"] = "无数据"
        return
    headers = list(rows[0].keys())
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append([row[header] for header in headers])
        if highlight_role:
            fill = KEPT_FILL if row.get("处理") == "保留" else REMOVED_FILL
            for cell in sheet[sheet.max_row]:
                cell.fill = fill
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
    for column_index, header in enumerate(headers, start=1):
        width = 18
        if header in {"标题", "链接", "被消除来源"}:
            width = 46
        elif header in {"判定依据", "域名"}:
            width = 28
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.row_dimensions[1].height = 22


def write_workbook(path: str, kept_rows: list[dict], detail_rows: list[dict], summary: dict, source: str) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "统计"
    summary_sheet.append(["项目", "数值"])
    for cell in summary_sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for key, value in summary.items():
        summary_sheet.append([key, value])
    summary_sheet.append(["数据目录", source])
    summary_sheet.append(["说明", "同一新闻保留正文更完整的一条；原始报告文件未改动。"])
    summary_sheet.column_dimensions["A"].width = 24
    summary_sheet.column_dimensions["B"].width = 88
    write_sheet(workbook, "去重结果", kept_rows)
    removed_rows = [row for row in detail_rows if row["处理"] == "消除" or row["组号"] in {
        item["组号"] for item in detail_rows if item["处理"] == "消除"
    }]
    write_sheet(workbook, "重复组明细", removed_rows, highlight_role=True)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Deduplicate same news stories into an Excel table.")
    parser.add_argument("--input", default=DEFAULT_INPUT, help="Reports directory to read.")
    parser.add_argument("--output", default=DEFAULT_OUTPUT, help="Excel path to write.")
    args = parser.parse_args()

    articles = load_articles(args.input)
    kept_rows, detail_rows, summary = build_rows(articles)
    write_workbook(args.output, kept_rows, detail_rows, summary, args.input)
    sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(args.output)


if __name__ == "__main__":
    main()
